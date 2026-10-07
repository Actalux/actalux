"""Tests for the document parser."""

from pathlib import Path
from textwrap import dedent

import pytest

from actalux.errors import ParseError
from actalux.ingest.parser import exotic_char_ratio, parse_file, strip_control_chars


class TestExoticCharRatio:
    def test_clean_english_is_zero(self) -> None:
        assert exotic_char_ratio("The board approved the FY2024 budget.") == 0.0

    def test_smart_punctuation_not_counted(self) -> None:
        # Curly quotes and em dashes are legitimate (General Punctuation block).
        assert exotic_char_ratio("the “levy” — passed 5-0") == 0.0

    def test_mojibake_scores_high(self) -> None:
        # Broken-font Cyrillic-block glyphs (the doc 87 cost tables) score high.
        assert exotic_char_ratio("ҨҨ ьы҂ эыэѐ") > 0.5

    def test_empty_is_zero(self) -> None:
        assert exotic_char_ratio("") == 0.0


@pytest.fixture
def tmp_dir(tmp_path: Path) -> Path:
    return tmp_path


class TestStripControlChars:
    def test_replaces_control_chars_with_space(self) -> None:
        # The 0x08 / 0x01 artifacts seen in extracted PDFs become spaces.
        assert strip_control_chars("Planning\x08\n3") == "Planning \n3"
        assert strip_control_chars("COST\x01\x14ESTIMATE") == "COST  ESTIMATE"

    def test_keeps_tab_newline_carriage_return(self) -> None:
        assert strip_control_chars("a\tb\nc\rd") == "a\tb\nc\rd"

    def test_strips_c1_controls(self) -> None:
        # C1 bytes (0x80-0x9f) from broken fonts are artifacts, not punctuation.
        assert strip_control_chars("bullet\x82 item") == "bullet  item"

    def test_clean_text_unchanged(self) -> None:
        text = "Ordinary minutes — approved 5-0. The “levy” passed."
        assert strip_control_chars(text) == text

    def test_applied_during_parse(self, tmp_dir: Path) -> None:
        md = tmp_dir / "dirty.md"
        md.write_bytes(b"# Heading\x08\n\nBody\x01text here.")
        result = parse_file(md)
        assert "\x08" not in result and "\x01" not in result
        assert "Body text here." in result


class TestParseMarkdown:
    def test_basic_markdown(self, tmp_dir: Path) -> None:
        md = tmp_dir / "minutes.md"
        md.write_text("# Board Meeting\n\nThe meeting was called to order.")
        result = parse_file(md)
        assert "Board Meeting" in result
        assert "called to order" in result

    def test_empty_file_raises(self, tmp_dir: Path) -> None:
        md = tmp_dir / "empty.md"
        md.write_text("")
        with pytest.raises(ParseError, match="empty"):
            parse_file(md)

    def test_whitespace_only_raises(self, tmp_dir: Path) -> None:
        md = tmp_dir / "blank.md"
        md.write_text("   \n\n   ")
        with pytest.raises(ParseError, match="empty"):
            parse_file(md)


class TestParseHtml:
    def test_basic_html(self, tmp_dir: Path) -> None:
        html = tmp_dir / "agenda.html"
        html.write_text(
            dedent("""\
            <html><body>
            <h1>Board Meeting Agenda</h1>
            <p>1. Call to Order</p>
            <p>2. Budget Discussion</p>
            </body></html>
        """)
        )
        result = parse_file(html)
        assert "Board Meeting Agenda" in result
        assert "Budget Discussion" in result

    def test_strips_script_and_style(self, tmp_dir: Path) -> None:
        html = tmp_dir / "messy.html"
        html.write_text(
            dedent("""\
            <html><body>
            <script>alert('xss')</script>
            <style>body { color: red; }</style>
            <p>Real content here.</p>
            </body></html>
        """)
        )
        result = parse_file(html)
        assert "alert" not in result
        assert "color: red" not in result
        assert "Real content here" in result


class TestParseText:
    def test_plain_text(self, tmp_dir: Path) -> None:
        txt = tmp_dir / "notes.txt"
        txt.write_text("Board discussed the proposed tax levy increase.")
        result = parse_file(txt)
        assert "tax levy" in result


class TestUnsupportedFormat:
    def test_unsupported_extension(self, tmp_dir: Path) -> None:
        doc = tmp_dir / "spreadsheet.xlsx"
        doc.write_bytes(b"not a real xlsx")
        with pytest.raises(ParseError, match="Unsupported"):
            parse_file(doc)


class TestParseDocx:
    def test_paragraphs_and_table_cells_in_document_order(self, tmp_path: Path) -> None:
        from docx import Document

        doc = Document()
        doc.add_paragraph("Board of Education Meeting - Oct 29 2025")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "1.1"
        table.rows[0].cells[1].text = "Call to Order"
        doc.add_paragraph("Moved by: Ms. Chris Win")
        path = tmp_path / "minutes.docx"
        doc.save(path)

        text = parse_file(path)
        assert text.splitlines() == [
            "Board of Education Meeting - Oct 29 2025",
            "1.1",
            "Call to Order",
            "Moved by: Ms. Chris Win",
        ]

    def test_merged_cells_are_not_repeated(self, tmp_path: Path) -> None:
        from docx import Document

        doc = Document()
        table = doc.add_table(rows=1, cols=3)
        merged = table.rows[0].cells[0].merge(table.rows[0].cells[1])
        merged.text = "Roll call"
        table.rows[0].cells[2].text = "Aye"
        path = tmp_path / "t.docx"
        doc.save(path)
        assert parse_file(path) == "Roll call\nAye"


class TestParsePptx:
    def test_slides_text_tables_and_notes(self, tmp_path: Path) -> None:
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        s1 = prs.slides.add_slide(prs.slide_layouts[1])
        s1.shapes.title.text = "Cognia Presentation"
        s1.notes_slide.notes_text_frame.text = "Speaker note"
        s2 = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
        tbl = s2.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
        tbl.cell(0, 0).text = "Standard"
        tbl.cell(0, 1).text = "Met"
        prs.slides.add_slide(prs.slide_layouts[6])  # image-only / empty slide
        path = tmp_path / "deck.pptx"
        prs.save(path)

        text = parse_file(path)
        assert "Slide 1\nCognia Presentation" in text
        assert "Speaker note" in text
        assert "Slide 2\nStandard | Met" in text
        assert "Slide 3" not in text


class TestScannedPdf:
    def test_image_only_pdf_is_ocrd_when_tesseract_returns_text(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        import fitz

        from actalux.ingest import parser

        doc = fitz.open()
        doc.new_page()  # a page with no text layer, like a scan
        path = tmp_path / "scan.pdf"
        doc.save(path)
        monkeypatch.setattr(parser, "_ocr_page", lambda page: "This Agreement is made")
        assert parse_file(path) == "This Agreement is made"

    def test_image_only_pdf_still_fails_without_ocr(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        import fitz

        from actalux.ingest import parser

        doc = fitz.open()
        doc.new_page()
        path = tmp_path / "scan.pdf"
        doc.save(path)
        monkeypatch.setattr(parser, "_ocr_page", lambda page: "")
        with pytest.raises(ParseError, match="no extractable text"):
            parse_file(path)
